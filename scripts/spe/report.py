#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
# SPDX-FileCopyrightText: Copyright contributors to the vLLM project

"""Export complete SPE CSV details and append one compact instruction sheet."""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import logging
import math
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from collections import Counter
from collections.abc import Iterable, Iterator
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import TextIO
from xml.sax.saxutils import escape, quoteattr
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment

LOGGER = logging.getLogger(__name__)
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SHEET_NAMES = ("SPE",)
MAX_ROWS = 1_048_576
COUNTERS = (("0", "TOT"), ("1", "ISSUE"), ("2", "XLAT"), ("6", "Counter 6"))
NS = {"s": MAIN_NS}
EVENT_NAMES = {
    0: "EXCEPTION-GEN",
    1: "RETIRED",
    2: "L1D-ACCESS",
    3: "L1D-REFILL",
    4: "TLB-ACCESS",
    5: "TLB-REFILL",
    6: "NOT-TAKEN",
    7: "MISPRED",
    8: "LLC-ACCESS",
    9: "LLC-REFILL",
    10: "REMOTE-ACCESS",
    11: "ALIGNMENT",
    17: "SVE-PARTIAL-PRED",
    18: "SVE-EMPTY-PRED",
}
CellValue = str | int | float | None


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    width: int = 18


@dataclass(frozen=True)
class Formula:
    expression: str
    cached: int | None


@dataclass(frozen=True)
class Styles:
    header: int
    body: int
    integer: int
    decimal: int


SAMPLE_COLUMNS = (
    Column("record_id", "Record ID", 30),
    Column("cpu", "CPU", 9),
    Column("tid", "主机 TID", 12),
    Column("container_tid", "容器 TID", 12),
    Column("request", "请求", 9),
    Column("step", "解码步", 9),
    Column("phase", "阶段", 23),
    Column("library", "代码库", 28),
    Column("function", "函数", 42),
    Column("instruction", "机器指令", 34),
    Column("pc", "运行时 PC", 23),
    Column("elf_pc", "ELF PC", 23),
    Column("function_start", "函数起点〔ELF〕", 23),
    Column("function_end", "函数终点〔ELF〕", 23),
    Column("va", "数据 VA", 23),
    Column("pa", "数据 PA", 23),
    Column("physical_resource", "物理地址资源", 22),
    Column("ram_numa", "RAM NUMA 节点", 20),
    Column("pci_bdf", "PCI BDF", 20),
    Column("bar_offset", "BAR 内偏移", 23),
    Column("data_mapping", "数据 VA 映射", 72),
    *(
        Column(
            name.lower().replace(" ", ""),
            f"{name}\n{'原值' if key == '6' else 'cycles'}",
            16,
        )
        for key, name in COUNTERS
    ),
    Column("difference", "TOT−ISSUE\n派生 cycles", 18),
    *(Column(f"present_{key}", f"{name}\n包存在", 18) for key, name in COUNTERS),
    Column("event_bits", "事件原始位图", 20),
    Column("event_width", "事件载荷位宽", 17),
    *(
        Column(
            f"event_{bit}", f"bit{bit} {EVENT_NAMES.get(bit, '语义未确认')}\n原始位", 22
        )
        for bit in range(32)
    ),
    Column("unknown_event_bits", "未命名事件位图", 22),
    Column("ticks", "CNTVCT ticks", 24),
    Column("time_ns", "时间戳 ns", 24),
    Column("operation", "原始操作类型", 24),
    Column("op_class", "OP 类别", 12),
    Column("op_value", "OP 原始载荷", 20),
    Column("data_source", "Data Source 原始载荷", 25),
    Column("context_id", "CONTEXT 原始 ID", 23),
    Column("context_el", "CONTEXT EL", 14),
    Column("counter_values", "全部 Counter 索引:原值", 48),
    Column("counter_presence", "全部 Counter 包存在", 48),
    Column("stream_id", "AUX 流", 25),
    Column("aux_tid", "AUX TID", 12),
    Column("aux_index", "AUX 序号", 12),
    Column("file_offset", "代码文件偏移", 23),
    Column("binary_sha256", "代码 SHA256", 68),
    Column("binary_build_id", "代码 Build ID", 44),
    Column("instruction_bytes", "指令机器码", 24),
    Column("resolution_status", "代码解析状态", 28),
    Column("function_status", "函数解析状态", 28),
    Column("instruction_status", "指令解析状态", 28),
    Column("pc_key", "指令分组标识", 70),
    Column("pc_payload", "PC 包原始载荷", 23),
    Column("va_payload", "VA 包原始载荷", 23),
    Column("pa_payload", "PA 包原始载荷", 23),
    Column("file", "AUX 来源文件", 38),
    Column("raw_hex", "样本原始 hex", 70),
    Column("raw_packets", "原始包序列", 80),
    Column("cache_location", "缓存供数位置〔未判定〕", 26),
    Column("memory_policy", "内存映射策略〔未判定〕", 26),
)

PC_COLUMNS = (
    Column("library", "代码库", 28),
    Column("function", "函数", 42),
    Column("instruction", "机器指令", 34),
    Column("samples", "样本数\n〔非执行次数〕", 24),
    Column("elf_pc", "ELF PC", 23),
    Column("pcs", "运行时 PC", 26),
    Column("function_start", "函数起点〔ELF〕", 23),
    Column("function_end", "函数终点〔ELF〕", 23),
    Column("file_offset", "代码文件偏移", 23),
    Column("physical_resource_counts", "物理资源:样本数", 32),
    Column("ram_numa_counts", "RAM NUMA 节点:样本数", 32),
    Column("pci_bdf_counts", "PCI BDF:样本数", 32),
    Column("data_mapping_counts", "数据 VA 映射:样本数", 72),
    Column("cpus", "CPU:样本数", 30),
    Column("tids", "主机 TID:样本数", 30),
    Column("requests", "请求:样本数", 30),
    Column("replay_samples", "Replay 内样本数", 20),
    Column("event_present", "含事件包样本数", 20),
    *(
        Column(
            f"event_{bit}",
            f"bit{bit} {EVENT_NAMES.get(bit, '语义未确认')}\n置1样本数",
            22,
        )
        for bit in range(32)
    ),
    *(
        Column(
            f"counter_{key}_{stat}",
            f"{name}\n{label.replace('cycles', '原值') if key == '6' else label}",
            19,
        )
        for key, name in COUNTERS
        for stat, label in (
            ("present", "有包样本数"),
            ("missing", "缺包样本数"),
            ("min", "最小 cycles"),
            ("max", "最大 cycles"),
            ("mean", "均值 cycles"),
            ("median", "中位数 cycles"),
            ("p95", "P95 cycles"),
        )
    ),
    Column("operation_counts", "操作类型:样本数", 40),
    Column("unknown_event_values", "未命名事件位图:样本数", 42),
    Column("counter_stats", "全部 Counter 存在与分位统计", 80),
    Column("binary_sha256", "代码 SHA256", 68),
    Column("binary_build_id", "代码 Build ID", 44),
    Column("instruction_bytes", "指令机器码", 24),
    Column("resolution_status", "代码解析状态", 28),
    Column("function_status", "函数解析状态", 28),
    Column("instruction_status", "指令解析状态", 28),
    Column("pc_key", "指令分组标识", 70),
)


SUMMARY_KEYS = (
    "library",
    "elf_pc",
    "pcs",
    "function",
    "instruction",
    "samples",
    "sample_share",
    "counter_0_mean",
    "counter_0_median",
    "counter_0_p95",
    "counter_1_mean",
    "counter_2_mean",
    "counter_2_p95",
    "event_2_rate",
    "event_3_rate",
    "event_4_rate",
    "event_5_rate",
    "event_8_rate",
    "event_9_rate",
    "event_10_rate",
    "ram_numa_counts",
    "physical_resource_counts",
)
SUMMARY_COLUMNS = tuple(
    next(
        (column for column in PC_COLUMNS if column.key == key),
        Column(
            key,
            "样本占比"
            if key == "sample_share"
            else f"{EVENT_NAMES.get(int(key.split('_')[1]), '')}\n占含事件包样本"
            if key.startswith("event_")
            else key,
            24,
        ),
    )
    for key in SUMMARY_KEYS
)


def summary_values(pc: dict, total: int) -> list[CellValue]:
    values = dict(
        zip((column.key for column in PC_COLUMNS), pc_values(pc), strict=True)
    )
    values["sample_share"] = pc["samples"] / total
    present = pc.get("event_present", 0)
    for bit in (2, 3, 4, 5, 8, 9, 10):
        values[f"event_{bit}_rate"] = (
            values[f"event_{bit}"] / present if present else None
        )
    return [values.get(key) for key in SUMMARY_KEYS]


def format_summary(workbook: Path) -> None:
    saved = load_workbook(workbook)
    sheet = saved["SPE"]
    template = saved.worksheets[0]
    for index, key in enumerate(SUMMARY_KEYS, 1):
        numeric = key in {"samples", "sample_share"} or key.startswith(
            ("counter_", "event_")
        )
        percentage = key == "sample_share" or key.endswith("_rate")
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, index)
            reference = template.cell(4, 2)
            cell.font = copy(reference.font)
            cell.fill = copy(reference.fill)
            cell.border = copy(reference.border)
            cell.alignment = Alignment(
                horizontal="right" if numeric else "left",
                vertical="center",
                indent=1,
            )
            cell.number_format = (
                "0.00%"
                if percentage
                else "#,##0"
                if key == "samples"
                else "0.00"
                if numeric
                else "@"
            )
        header = sheet.cell(1, index)
        reference = template.cell(1, 2)
        header.font = copy(reference.font)
        header.fill = copy(reference.fill)
        header.border = copy(reference.border)
        alignment = copy(header.alignment)
        alignment.horizontal = "center"
        alignment.indent = 1
        header.alignment = alignment
    sheet.freeze_panes = "D2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 26
    sheet.column_dimensions["D"].width = 38
    sheet.column_dimensions["E"].width = 34
    saved.save(workbook)
    saved.close()


def export_csv(path: Path, columns: tuple[Column, ...], rows: Iterable[list]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(column.key for column in columns)
        for row in rows:
            writer.writerow(
                value.cached if isinstance(value, Formula) else value for value in row
            )


def mapping(value: object, context: str) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{context} must be an object")
    return value


def compact(value: object) -> str | None:
    if value is None:
        return None
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def sample_rows(path: Path) -> Iterator[dict]:
    with gzip.open(path, "rt", encoding="utf-8") as stream:
        for line_number, line in enumerate(stream, 1):
            if line.strip():
                yield mapping(json.loads(line), f"sample line {line_number}")


def counter_values(sample: dict) -> dict:
    counters = mapping(sample.get("counters"), "counters")
    presence = mapping(sample.get("counter_presence"), "counter_presence")
    for index in range(32):
        key = str(index)
        if presence.get(key) is not (key in counters):
            raise ValueError(f"Counter {key} presence disagrees with raw counters")
    for key, value in counters.items():
        if not key.isdecimal() or not 0 <= int(key) < 32:
            raise ValueError(f"invalid counter index {key}")
        if not isinstance(value, int) or isinstance(value, bool) or value < 0:
            raise ValueError(f"invalid Counter {key} value")
    return counters


def sample_values(sample: dict, row_number: int) -> list[CellValue | Formula]:
    counters = counter_values(sample)
    values = dict(sample)
    values["library"] = Path(sample["binary"]).name if sample.get("binary") else None
    values["phase"] = "Replay" if sample.get("replay") else "run_fullgraph 其余"
    values["counter_values"] = compact(counters)
    values["counter_presence"] = compact(sample["counter_presence"])
    values["raw_packets"] = compact(sample.get("raw_packets"))
    values["cache_location"] = "unknown"
    values["memory_policy"] = "unknown"
    event_bits = sample.get("event_bits")
    events = int(event_bits, 16) if event_bits is not None else None
    if events is not None and not 0 <= events < 2**64:
        raise ValueError("event bitmap exceeds the 64-bit packet payload")
    for bit in range(32):
        values[f"event_{bit}"] = (events >> bit) & 1 if events is not None else None
    for key, name in COUNTERS:
        field = name.lower().replace(" ", "")
        values[field] = counters.get(key)
        values[f"present_{key}"] = int(key in counters)
        if sample.get(field) != counters.get(key):
            raise ValueError(f"{field} disagrees with raw counters")
    keys = [column.key for column in SAMPLE_COLUMNS]
    tot_ref = f"{column_letter(keys.index('tot') + 1)}{row_number}"
    issue_ref = f"{column_letter(keys.index('issue') + 1)}{row_number}"
    difference = (
        counters["0"] - counters["1"] if {"0", "1"} <= counters.keys() else None
    )
    values["difference"] = Formula(
        f'IF(COUNT({tot_ref},{issue_ref})=2,{tot_ref}-{issue_ref},"")', difference
    )
    for key in (
        "ticks",
        "time_ns",
        "pc",
        "va",
        "pa",
        "pc_payload",
        "va_payload",
        "pa_payload",
        "elf_pc",
        "file_offset",
        "bar_offset",
        "function_start",
        "function_end",
    ):
        if sample.get(key) is not None and not isinstance(sample[key], str):
            raise ValueError(f"{key} must be exact text")
    return [values.get(column.key) for column in SAMPLE_COLUMNS]


def pc_values(pc: dict) -> list[CellValue]:
    values = dict(pc)
    values["library"] = Path(pc["binary"]).name if pc.get("binary") else None
    for key in (
        "cpus",
        "tids",
        "requests",
        "operation_counts",
        "unknown_event_values",
        "physical_resource_counts",
        "ram_numa_counts",
        "pci_bdf_counts",
        "data_mapping_counts",
    ):
        values[key] = compact(pc.get(key))
    values["pcs"] = ", ".join(pc.get("pcs", []))
    counters = mapping(pc.get("counters"), "PC counters")
    values["counter_stats"] = compact(counters)
    events = mapping(pc.get("event_bits"), "PC event counts")
    for bit in range(32):
        values[f"event_{bit}"] = events.get(str(bit), 0)
    for key, _ in COUNTERS:
        stats = mapping(counters.get(key), f"Counter {key} statistics")
        if stats.get("present", 0) + stats.get("missing", 0) != pc["samples"]:
            raise ValueError(f"Counter {key} population differs from PC samples")
        for stat in ("present", "missing", "min", "max", "mean", "median", "p95"):
            values[f"counter_{key}_{stat}"] = stats.get(stat)
    return [values.get(column.key) for column in PC_COLUMNS]


def inherited_styles(source: ZipFile) -> Styles:
    root = ET.fromstring(source.read("xl/styles.xml"))
    fonts = root.find("s:fonts", NS)
    fills = root.find("s:fills", NS)
    xfs = root.find("s:cellXfs", NS)
    if fonts is None or fills is None or xfs is None:
        raise ValueError("Topdown workbook style table is incomplete")
    matches = {}
    for index, xf in enumerate(xfs):
        font = fonts[int(xf.get("fontId", "0"))]
        name = font.find("s:name", NS)
        if name is None or name.get("val") not in {"Carlito", "微软雅黑"}:
            continue
        fill = fills[int(xf.get("fillId", "0"))].find("s:patternFill/s:fgColor", NS)
        color = fill.get("rgb") if fill is not None else None
        if color == "FF1F4E78" and font.find("s:b", NS) is not None:
            matches["header"] = index
        elif color in {None, "FFFFFFFF"} and font.find("s:b", NS) is None:
            key = {"0": "body", "1": "integer", "2": "decimal"}.get(xf.get("numFmtId"))
            if key:
                matches[key] = index
    if "header" not in matches or "body" not in matches:
        raise ValueError("expected existing Topdown detail styles were not found")
    return Styles(
        matches["header"],
        matches["body"],
        matches.get("integer", matches["body"]),
        matches.get("decimal", matches["body"]),
    )


def write_cell(
    stream: TextIO,
    reference: str,
    value: CellValue | Formula,
    styles: Styles,
    header: bool = False,
) -> None:
    if isinstance(value, Formula):
        kind = ' t="str"' if value.cached is None else ""
        cached = "" if value.cached is None else str(value.cached)
        stream.write(
            f'<c r="{reference}" s="{styles.integer}"{kind}>'
            f"<f>{escape(value.expression)}</f><v>{cached}</v></c>"
        )
    elif value is None:
        stream.write(f'<c r="{reference}" s="{styles.body}"/>')
    elif isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"nonfinite cell {reference}")
        style = styles.integer if isinstance(value, int) else styles.decimal
        stream.write(f'<c r="{reference}" s="{style}"><v>{value}</v></c>')
    elif isinstance(value, str):
        if len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
            raise ValueError(
                f"cell {reference} cannot be represented losslessly in Excel"
            )
        style = styles.header if header else styles.body
        stream.write(
            f'<c r="{reference}" s="{style}" t="inlineStr">'
            f'<is><t xml:space="preserve">{escape(value)}</t></is></c>'
        )
    else:
        raise ValueError(f"unsupported cell value at {reference}")


def write_sheet(
    path: Path,
    columns: tuple[Column, ...],
    rows: Iterable[list[CellValue | Formula]],
    count: int,
    styles: Styles,
) -> None:
    if not 0 < count < MAX_ROWS:
        raise ValueError("SPE row count must be nonzero and fit one Excel worksheet")
    last = f"{column_letter(len(columns))}{count + 1}"
    with path.open("w", encoding="utf-8") as stream:
        stream.write(
            f'<worksheet xmlns="{MAIN_NS}"><dimension ref="A1:{last}"/>'
            '<sheetViews><sheetView showGridLines="0" workbookViewId="0">'
            '<pane xSplit="1" ySplit="1" topLeftCell="B2" '
            'activePane="bottomRight" state="frozen"/>'
            '<selection pane="bottomRight" activeCell="B2" sqref="B2"/>'
            '</sheetView></sheetViews><sheetFormatPr defaultRowHeight="21"/>'
            "<cols>"
        )
        for index, column in enumerate(columns, 1):
            stream.write(
                f'<col min="{index}" max="{index}" width="{column.width}" '
                'customWidth="1"/>'
            )
        stream.write('</cols><sheetData><row r="1" ht="42" customHeight="1">')
        for index, column in enumerate(columns, 1):
            write_cell(stream, f"{column_letter(index)}1", column.label, styles, True)
        stream.write("</row>")
        actual = 0
        for row_number, row in enumerate(rows, 2):
            if len(row) != len(columns):
                raise ValueError("SPE row width differs from its headers")
            actual += 1
            stream.write(f'<row r="{row_number}" ht="21" customHeight="1">')
            for index, value in enumerate(row, 1):
                write_cell(stream, f"{column_letter(index)}{row_number}", value, styles)
            stream.write("</row>")
        if actual != count:
            raise ValueError("SPE sample count changed during report generation")
        stream.write(
            f'</sheetData><autoFilter ref="A1:{last}"/>'
            '<pageMargins left="0.75" right="0.75" top="1" bottom="1" '
            'header="0.5" footer="0.5"/></worksheet>'
        )


def append_xml(data: bytes, tag: str, addition: str) -> bytes:
    closing = f"</{tag}>".encode()
    if data.count(closing) != 1:
        raise ValueError(f"unsupported workbook XML container {tag}")
    return data.replace(closing, addition.encode() + closing, 1)


def append_package(
    source_path: Path, output_path: Path, sheet_paths: list[Path]
) -> None:
    with ZipFile(source_path) as source:
        workbook_bytes = source.read("xl/workbook.xml")
        workbook = ET.fromstring(workbook_bytes)
        sheets = workbook.find("s:sheets", NS)
        if sheets is None:
            raise ValueError("workbook has no sheets")
        if any(sheet.get("name") in SHEET_NAMES for sheet in sheets):
            raise ValueError(
                "SPE sheets already exist; rebuild the Topdown workbook first"
            )
        next_id = max(int(sheet.get("sheetId", "0")) for sheet in sheets) + 1
        relationships = source.read("xl/_rels/workbook.xml.rels")
        rel_root = ET.fromstring(relationships)
        used = {rel.get("Id") for rel in rel_root}
        added_sheets = []
        added_relationships = []
        added_types = []
        entries = []
        for offset, (name, path) in enumerate(
            zip(SHEET_NAMES, sheet_paths, strict=True)
        ):
            sheet_id = next_id + offset
            entry = f"xl/worksheets/spe{sheet_id}.xml"
            if entry in source.namelist():
                raise ValueError(f"worksheet package entry already exists: {entry}")
            rid = f"rIdSPE{sheet_id}"
            if rid in used:
                raise ValueError("SPE worksheet relationship already exists")
            added_sheets.append(
                f'<sheet xmlns:r="{REL_NS}" name={quoteattr(name)} '
                f'sheetId="{sheet_id}" r:id="{rid}"/>'
            )
            added_relationships.append(
                f'<Relationship Id="{rid}" Type="{REL_NS}/worksheet" '
                f'Target="/xl/worksheets/spe{sheet_id}.xml"/>'
            )
            added_types.append(
                f'<Override PartName="/{entry}" '
                'ContentType="application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.worksheet+xml"/>'
            )
            entries.append((entry, path))
        replacements = {
            "xl/workbook.xml": append_xml(
                workbook_bytes, "sheets", "".join(added_sheets)
            ),
            "xl/_rels/workbook.xml.rels": append_xml(
                relationships, "Relationships", "".join(added_relationships)
            ),
            "[Content_Types].xml": append_xml(
                source.read("[Content_Types].xml"), "Types", "".join(added_types)
            ),
        }
        with ZipFile(output_path, "w", compression=ZIP_DEFLATED) as output:
            for info in source.infolist():
                if info.filename in replacements:
                    output.writestr(info, replacements[info.filename])
                else:
                    with source.open(info) as reader, output.open(info, "w") as writer:
                        shutil.copyfileobj(reader, writer, 1024 * 1024)
            for entry, path in entries:
                output.write(path, entry)


def append_report(workbook: Path, spe_dir: Path) -> dict[str, int]:
    """Validate decoded input, export details, and append the SPE summary."""
    audit = mapping(json.loads((spe_dir / "audit.json").read_text()), "SPE audit")
    if audit.get("status") != "pass":
        raise ValueError("SPE decoder audit did not pass")
    manifest = mapping(
        json.loads((spe_dir / "manifest.json").read_text()), "SPE manifest"
    )
    if manifest.get("status") != "pass":
        raise ValueError("SPE decoder manifest did not pass")
    pcs = json.loads((spe_dir / "pc_summary.json").read_text())
    if not isinstance(pcs, list) or not pcs:
        raise ValueError("SPE PC summary must contain at least one instruction")
    expected = Counter()
    for pc in pcs:
        pc = mapping(pc, "PC summary row")
        key = pc.get("pc_key")
        count = pc.get("samples")
        if not isinstance(key, str) or not key or key in expected:
            raise ValueError("PC group identifiers must be nonempty and unique")
        if not isinstance(count, int) or count <= 0:
            raise ValueError("PC sample counts must be positive integers")
        expected[key] = count
    sample_path = spe_dir / "samples.jsonl.gz"
    observed = Counter(sample.get("pc_key") for sample in sample_rows(sample_path))
    if observed != expected:
        raise ValueError("SPE samples and PC populations disagree")
    count = sum(observed.values())
    for metadata in (audit, manifest):
        if metadata.get("selected_samples") != count or metadata.get("pc_count") != len(
            pcs
        ):
            raise ValueError("SPE metadata totals disagree with decoded records")
    with tempfile.TemporaryDirectory(
        prefix=".spe_report_", dir=workbook.parent
    ) as name:
        temporary = Path(name)
        with ZipFile(workbook) as source:
            styles = inherited_styles(source)
        details = workbook.parent / "details"
        details.mkdir(exist_ok=True)
        export_csv(
            details / "spe_samples.csv",
            SAMPLE_COLUMNS,
            (
                sample_values(row, index)
                for index, row in enumerate(sample_rows(sample_path), 2)
            ),
        )
        export_csv(
            details / "spe_instructions.csv", PC_COLUMNS, (pc_values(pc) for pc in pcs)
        )
        pcs_sheet = temporary / "pcs.xml"
        ordered = sorted(pcs, key=lambda pc: (-pc["samples"], pc["pc_key"]))
        write_sheet(
            pcs_sheet,
            SUMMARY_COLUMNS,
            (summary_values(pc, count) for pc in ordered),
            len(pcs),
            styles,
        )
        output = temporary / "report.xlsx"
        append_package(workbook, output, [pcs_sheet])
        with ZipFile(output) as saved:
            if saved.testzip() is not None:
                raise ValueError("SPE report archive verification failed")
        format_summary(output)
        output.replace(workbook)
    return {"samples": count, "pcs": len(pcs)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--spe-dir", required=True, type=Path)
    args = parser.parse_args()
    totals = append_report(args.workbook, args.spe_dir)
    LOGGER.info(
        "appended SPE samples=%d PCs=%d to %s",
        totals["samples"],
        totals["pcs"],
        args.workbook,
    )
    return 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    raise SystemExit(main())
