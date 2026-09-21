# SPDX-License-Identifier: Apache-2.0
# SPDX-FileCopyrightText: Copyright contributors to the vLLM project

"""Compact report contracts: traceability, event denominators, structured hotspots."""

import csv
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from scripts import build_xlsx
from scripts.spe import report


def test_details_preserve_source_bytes_without_recursive_copy(tmp_path: Path) -> None:
    source = tmp_path / "topdown/raw/stage.csv"
    source.parent.mkdir(parents=True)
    source.write_text("sequence,cycles\n1,987654321\n")
    build_xlsx.export_details(tmp_path)
    build_xlsx.export_details(tmp_path)
    entries = json.loads((tmp_path / "details/index.json").read_text())
    assert len(entries) == 1
    assert (tmp_path / entries[0]["csv"]).read_bytes() == source.read_bytes()


def test_hotspot_preserves_symbols_and_fraction(tmp_path: Path) -> None:
    source = tmp_path / "perf_report.txt"
    source.write_text("# header\n  60.46%  Worker  lib.so  [.] foo(int, long)\n")
    sheet = Workbook().active
    build_xlsx.write_compact_hotspot(sheet, source, tmp_path)
    assert sheet.cell(2, 1).value == pytest.approx(0.6046)
    assert sheet.cell(2, 4).value == "[.] foo(int, long)"
    with (tmp_path / "details/hotspot.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.reader(stream))
    assert rows[1][3] == "[.] foo(int, long)"


def test_spe_csv_materializes_derived_value_and_preserves_missing(
    tmp_path: Path,
) -> None:
    path = tmp_path / "samples.csv"
    columns = (
        report.Column("difference", "difference", 12),
        report.Column("missing", "missing", 12),
    )
    report.export_csv(path, columns, [[report.Formula("A2-B2", 42), None]])
    with path.open(encoding="utf-8-sig") as stream:
        assert list(csv.reader(stream)) == [["difference", "missing"], ["42", ""]]


def test_spe_percentages_are_fractions_and_identity_stays_in_csv() -> None:
    pc = {
        "samples": 2,
        "event_present": 1,
        "event_bits": {"2": 1},
        "counters": {key: {"present": 0, "missing": 2} for key, _ in report.COUNTERS},
        "elf_pc": "0x1234",
        "pcs": ["0xabcd"],
        "pc_key": "binary-sha:0x1234",
    }
    values = dict(zip(report.SUMMARY_KEYS, report.summary_values(pc, 10), strict=True))
    assert values["sample_share"] == 0.2
    assert values["event_2_rate"] == 1
    assert report.SUMMARY_KEYS[:3] == ("library", "elf_pc", "pcs")
    assert "pc_key" not in report.SUMMARY_KEYS
    assert "pc_key" in {column.key for column in report.PC_COLUMNS}
    pc["event_present"] = 0
    values = dict(zip(report.SUMMARY_KEYS, report.summary_values(pc, 10), strict=True))
    assert values["event_2_rate"] is None
